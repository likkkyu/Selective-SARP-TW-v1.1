#!/usr/bin/env python3
"""Append a rigorous mathematical modeling slide to existing report deck."""

import argparse
import os
import shutil
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor


def add_textbox(slide, left, top, width, height, text, size=14, bold=False, color=(0, 0, 0), line_spacing=1.2):
    box = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height))
    tf = box.text_frame
    tf.clear()
    p = tf.paragraphs[0]
    run = p.add_run()
    run.text = text
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = RGBColor(*color)
    p.line_spacing = line_spacing
    return box


def add_bullets(slide, left, top, width, height, lines, size=13, color=(40, 40, 40), bullet='• '):
    box = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height))
    tf = box.text_frame
    tf.clear()
    for i, line in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = f"{bullet}{line}"
        p.level = 0
        p.line_spacing = 1.2
        if p.runs:
            p.runs[0].font.size = Pt(size)
            p.runs[0].font.color.rgb = RGBColor(*color)
    return box


def add_modeling_slide(input_ppt, output_ppt):
    if not os.path.exists(input_ppt):
        raise FileNotFoundError(f"Input PPT not found: {input_ppt}")

    os.makedirs(os.path.dirname(output_ppt), exist_ok=True)
    shutil.copy2(input_ppt, output_ppt)

    prs = Presentation(output_ppt)
    layout = prs.slide_layouts[6] if len(prs.slide_layouts) > 6 else prs.slide_layouts[-1]  # blank-like
    slide = prs.slides.add_slide(layout)

    # Header style block
    slide.shapes.add_shape(
        1,  # MSO_SHAPE.RECTANGLE
        Inches(0.0), Inches(0.0), Inches(13.333), Inches(0.58)
    ).fill.solid()
    # Need to re-fetch shape to set style robustly
    hdr = slide.shapes[-1]
    hdr.fill.solid()
    hdr.fill.fore_color.rgb = RGBColor(220, 235, 247)
    hdr.line.fill.background()

    add_textbox(
        slide,
        left=0.22,
        top=0.08,
        width=9.8,
        height=0.42,
        text="数学建模与问题定义（代码对齐）",
        size=24,
        bold=True,
        color=(31, 78, 121),
    )

    # Section A: problem definition
    slide.shapes.add_shape(1, Inches(0.35), Inches(0.75), Inches(12.6), Inches(1.35))
    sec_a = slide.shapes[-1]
    sec_a.fill.solid()
    sec_a.fill.fore_color.rgb = RGBColor(247, 249, 251)
    sec_a.line.color.rgb = RGBColor(210, 220, 230)

    add_textbox(slide, 0.5, 0.82, 3.5, 0.3, "1) 问题定义", size=16, bold=True, color=(31, 78, 121))
    add_bullets(
        slide,
        0.55,
        1.12,
        12.0,
        0.9,
        [
            "给定订单集合（乘客+货物）与固定运营窗 10:00–16:00，联合优化车辆-订单调度。",
            "动作空间包含节点访问与 reject 动作；目标是在硬约束下最小化 RMB 总成本。",
            "节点集：depot, pickups P, deliveries D；每个订单 i 满足 pickup-delivery 成对耦合。",
        ],
        size=12,
    )

    # Section B: objective function
    slide.shapes.add_shape(1, Inches(0.35), Inches(2.28), Inches(12.6), Inches(2.25))
    sec_b = slide.shapes[-1]
    sec_b.fill.solid()
    sec_b.fill.fore_color.rgb = RGBColor(255, 255, 255)
    sec_b.line.color.rgb = RGBColor(210, 220, 230)

    add_textbox(slide, 0.5, 2.35, 4.8, 0.3, "2) 评估目标函数（raw cost）", size=16, bold=True, color=(31, 78, 121))

    obj_formula = (
        "min J_raw = C_energy + C_p_delay + C_c_delay + C_vehicle + C_reject + C_tripOT\n"
        "C_energy = p_e Σ_t η(W_t)·d_t,   η(W)=0.18·(1+W/10000) [kWh/km], p_e=1.0\n"
        "C_p_delay = 0.6 Σ_{i∈D^P} δ_i^del(min),   C_c_delay = 0.06 Σ_{i∈D^C} δ_i^del(min)\n"
        "C_vehicle = 20·K_used,   C_reject = 500·N_rej,   C_tripOT = 200 Σ_k (T_k^trip-3)^+"
    )
    add_textbox(slide, 0.6, 2.72, 12.1, 1.72, obj_formula, size=12, color=(20, 20, 20), line_spacing=1.22)

    # Section C: hard constraints
    slide.shapes.add_shape(1, Inches(0.35), Inches(4.66), Inches(12.6), Inches(2.2))
    sec_c = slide.shapes[-1]
    sec_c.fill.solid()
    sec_c.fill.fore_color.rgb = RGBColor(247, 249, 251)
    sec_c.line.color.rgb = RGBColor(210, 220, 230)

    add_textbox(slide, 0.5, 4.73, 5.2, 0.3, "3) 核心硬约束（由 state mask 强制）", size=16, bold=True, color=(31, 78, 121))

    cons_lines = [
        "先序约束：delivery_i 仅在 pickup_i 完成后可访问；reject 后该订单 P/D 永久禁用。",
        "双舱容量：q_t^P ≤ Q^P, q_t^C ≤ Q^C；且载货时禁止提前回仓。",
        "乘客 pickup 硬时间窗：a_{p_i} ≤ b_{p_i}；乘客 ride-time 两级硬约束：ride_i≤70min, excess_i≤30min。",
        "行程/运营约束：单趟时长 ≤ 3h，且候选动作需满足回仓后不超过 16:00。",
        "车队上限：K_used ≤ ceil(N/6)。reject 动作仅对当前可行候选开放（确定性目标选择）。",
    ]
    add_bullets(slide, 0.58, 5.1, 12.15, 1.62, cons_lines, size=11)

    # footer source
    add_textbox(
        slide,
        0.45,
        6.95,
        12.3,
        0.28,
        "Source: problem_mcvrptw_v2.py / state_mcvrptw_v2.py / MODELING.md (current repository)",
        size=10,
        color=(90, 90, 90),
    )

    prs.save(output_ppt)
    print(f"[done] wrote PPT: {output_ppt}")


def parse_args():
    p = argparse.ArgumentParser(description="Append mathematical modeling slide")
    p.add_argument(
        "--input-ppt",
        default="/Users/bytedance/Downloads/2026年7月汇报_Selective_SARP_TW_阶段结果_v2_含参数表.pptx",
    )
    p.add_argument(
        "--output-ppt",
        default="/Users/bytedance/Downloads/2026年7月汇报_Selective_SARP_TW_阶段结果_v3_含参数表_含数学建模页.pptx",
    )
    return p.parse_args()


if __name__ == "__main__":
    args = parse_args()
    add_modeling_slide(args.input_ppt, args.output_ppt)
